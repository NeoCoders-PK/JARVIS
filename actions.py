"""
actions.py - The small set of things JARVIS is allowed to actually do on
your computer: open apps/URLs, and read/list/write files (for "review this
document" / "edit this document").

Every one of these is only ever called after main.py has shown you a plain-
English description of exactly what's about to happen and you've clicked
Yes. Nothing here runs unattended - it's called *by* the confirmation flow,
never bypassing it.

Each entry in ACTIONS has:
  - "declaration": the OpenAI-style function-calling schema (tells the
    model the function exists, what it does, and what arguments it takes -
    this is the format Groq's chat completions API expects)
  - "confirm_text": builds the human-readable confirmation prompt from args
  - "run": actually performs the action, returns a small dict result

If you want to add a new capability later (take a screenshot, run a shell
command, etc.), add one more entry here in the same shape - nothing else
needs to change.
"""
import csv
import io
import json
import os
import subprocess
import sys
import webbrowser


def _open_url(args):
    url = args.get("url", "")
    ok = webbrowser.open(url)
    if ok:
        return {"result": f"Opened {url} in the default browser."}
    return {"error": f"Could not open {url} - no browser handler found."}


def _open_app(args):
    path = args.get("path", "")
    try:
        if sys.platform == "win32":
            os.startfile(path)  # works for .exe, documents, and folders alike
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
        return {"result": f"Launched {path}."}
    except Exception as e:
        return {"error": f"Couldn't open {path}: {e}"}


def _list_directory(args):
    path = args.get("path", ".")
    try:
        entries = os.listdir(path)
        return {"result": entries[:200]}  # cap so a huge folder can't blow up context
    except Exception as e:
        return {"error": f"Couldn't list {path}: {e}"}


MAX_READ_CHARS = 20000  # keep huge files from blowing up the request


def _ext(path):
    return os.path.splitext(path)[1].lower()


# ---- format-specific read/write --------------------------------------------
# Each of these turns a rich document into/from a single plain-text string,
# using a documented convention (see SYSTEM_PROMPT in main.py, which tells
# the model these conventions). None of this preserves the original file's
# formatting, images, or styling - editing means "extract text, let the
# model change it, write a brand-new document with that text," not
# in-place modification of the original layout.

def _read_docx(path):
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def _write_docx(path, content):
    from docx import Document
    doc = Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    doc.save(path)


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    sheet = wb.active  # only the active/first sheet - see docstring above
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(",".join("" if c is None else str(c) for c in row))
    return "\n".join(rows)


def _write_xlsx(path, content):
    from openpyxl import Workbook
    wb = Workbook()
    sheet = wb.active
    reader = csv.reader(io.StringIO(content))
    for row in reader:
        sheet.append(row)
    wb.save(path)


def _read_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _write_pdf(path, content):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(path, pagesize=LETTER)
    width, height = LETTER
    margin, line_height = 54, 14
    y = height - margin
    for line in content.split("\n"):
        if y < margin:
            c.showPage()
            y = height - margin
        c.drawString(margin, y, line[:110])  # crude line-length cap, no wrapping
        y -= line_height
    c.save()


def _read_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    slide_blocks = []
    for slide in prs.slides:
        lines = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if text.strip():
                    lines.append(text)
        slide_blocks.append("\n".join(lines))
    return "\n---\n".join(slide_blocks)


def _strip_bullet_prefix(line):
    line = line.strip()
    for prefix in ("- ", "* ", "• "):
        if line.startswith(prefix):
            return line[len(prefix):]
    return line


def _write_pptx(path, content):
    from pptx import Presentation
    prs = Presentation()
    layout = prs.slide_layouts[1]  # "Title and Content"
    for block in content.split("---"):
        lines = [l for l in block.strip().split("\n") if l.strip()]
        if not lines:
            continue
        title = lines[0]
        bullets = [_strip_bullet_prefix(l) for l in lines[1:]]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        if bullets:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for b in bullets[1:]:
                body.add_paragraph().text = b
    prs.save(path)


def _create_presentation(args):
    """
    Builds a designed deck via presentation.py. Kept separate from
    write_file's plain .pptx path: that one takes loose text and makes
    basic title+bullet slides on PowerPoint's default template; this one
    takes a structured spec and produces a themed deck with real layouts,
    charts, web images, and transitions.
    """
    path = args.get("path", "")
    spec = args.get("spec", "")
    try:
        import presentation
    except ImportError as e:
        return {"error": f"Couldn't load the presentation module: {e}. "
                         "Run: pip install python-pptx lxml"}
    try:
        parsed = json.loads(spec) if isinstance(spec, str) else spec
    except json.JSONDecodeError as e:
        return {"error": f"The presentation spec isn't valid JSON: {e}"}

    # Pexels key is optional - without it, image slides fall back to
    # styled placeholder panels instead of failing.
    api_key = None
    try:
        cfg_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")
        with open(cfg_path, "r", encoding="utf-8") as f:
            api_key = json.load(f).get("pexels_api_key") or None
    except Exception:
        pass

    try:
        result = presentation.build_presentation(path, parsed, api_key)
    except Exception as e:
        return {"error": f"Couldn't build the presentation: {e!r}"}

    if not api_key:
        result["note"] = ("No pexels_api_key in config.json, so any image slides "
                          "used placeholder panels. Add a free key from "
                          "pexels.com/api to pull real photos.")
    return result


def _presentation_confirm_text(args):
    """Renders a readable slide-by-slide outline for the confirm dialog -
    raw JSON would be unreadable in a Yes/No popup."""
    path = args.get("path", "")
    # Deliberately NOT _write_warning(): that one warns the .pptx will be a
    # plain default-template deck, which is true for write_file but the
    # opposite of what this action does. The only caveat worth raising here
    # is the overwrite.
    warning = ""
    if path and os.path.exists(path):
        warning = "⚠️ This file already exists and will be overwritten.\n\n"
    try:
        spec = args.get("spec", "")
        parsed = json.loads(spec) if isinstance(spec, str) else spec
        slides = parsed.get("slides", [])
    except Exception:
        return (f"{warning}Create this presentation?\n\n{path}\n\n"
                "(couldn't preview the outline - the spec isn't valid JSON)")

    lines = []
    for i, s in enumerate(slides, 1):
        layout = s.get("layout", "bullets")
        title = s.get("title") or s.get("quote") or s.get("heading") or ""
        img = "  [+photo]" if s.get("image") else ""
        lines.append(f"{i}. [{layout}]{img} {title}"[:110])

    transition = parsed.get("transition", "fade")
    return (
        f"{warning}Create this presentation?\n\n{path}\n\n"
        f"{len(slides)} slides, '{transition}' transitions\n\n"
        + "\n".join(lines)
    )


def _read_file(args):
    path = args.get("path", "")
    ext = _ext(path)
    try:
        if ext == ".docx":
            content = _read_docx(path)
        elif ext == ".xlsx":
            content = _read_xlsx(path)
        elif ext == ".pdf":
            content = _read_pdf(path)
        elif ext == ".pptx":
            content = _read_pptx(path)
        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read(MAX_READ_CHARS + 1)
        return {"result": content[:MAX_READ_CHARS]}
    except ImportError as e:
        return {"error": f"Missing library to read {ext} files: {e}. Run: pip install python-docx openpyxl pypdf python-pptx"}
    except Exception as e:
        return {"error": f"Couldn't read {path}: {e}"}


def _write_file(args):
    path = args.get("path", "")
    content = args.get("content", "")
    ext = _ext(path)
    try:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        if ext == ".docx":
            _write_docx(path, content)
        elif ext == ".xlsx":
            _write_xlsx(path, content)
        elif ext == ".pdf":
            _write_pdf(path, content)
        elif ext == ".pptx":
            _write_pptx(path, content)
        else:
            with open(path, "w", encoding="utf-8") as f:
                f.write(content)
        return {"result": f"Wrote {len(content)} characters to {path}."}
    except ImportError as e:
        return {"error": f"Missing library to write {ext} files: {e}. Run: pip install python-docx openpyxl reportlab python-pptx"}
    except Exception as e:
        return {"error": f"Couldn't write {path}: {e}"}


def _write_warning(path):
    ext = _ext(path)
    if ext == ".docx":
        return "⚠️ This creates a plain new Word doc - original formatting/images will NOT carry over.\n\n"
    if ext == ".xlsx":
        return "⚠️ This creates a new single-sheet workbook from CSV-style rows - other sheets/formatting will NOT carry over.\n\n"
    if ext == ".pdf":
        return "⚠️ This creates a plain new PDF from text - original layout/fonts/images will NOT carry over.\n\n"
    if ext == ".pptx":
        return "⚠️ This creates a plain new deck (default template, title+bullets only) - no custom design/images/animations.\n\n"
    return ""


ACTIONS = {
    "open_url": {
        "declaration": {
            "name": "open_url",
            "description": "Open a web page in the user's default browser.",
            "parameters": {
                "type": "object",
                "properties": {"url": {"type": "string", "description": "The full URL to open."}},
                "required": ["url"],
            },
        },
        "confirm_text": lambda a: f"Open this URL in your browser?\n\n{a.get('url')}",
        "run": _open_url,
    },
    "open_app": {
        "declaration": {
            "name": "open_app",
            "description": "Launch an application or open a file/document with its default program.",
            "parameters": {
                "type": "object",
                "properties": {"path": {"type": "string", "description": "Path to the app, executable, or file to open."}},
                "required": ["path"],
            },
        },
        "confirm_text": lambda a: f"Open this application or file?\n\n{a.get('path')}",
        "run": _open_app,
    },
    "list_directory": {
        "declaration": {
            "name": "list_directory",
            "description": "List the files and folders inside a directory, for reviewing what's there.",
            "parameters": {
                "type": "object",
                "properties": {"path": {"type": "string", "description": "Folder path to list."}},
                "required": ["path"],
            },
        },
        "confirm_text": lambda a: f"List the contents of this folder?\n\n{a.get('path')}",
        "run": _list_directory,
    },
    "read_file": {
        "declaration": {
            "name": "read_file",
            "description": (
                "Read and return the text content of a file, so it can be reviewed or "
                "summarized. Supports plain text, .docx (Word), .xlsx (Excel, active "
                "sheet only, returned as comma-separated rows), .pdf (extracted text "
                "only, formatting/images lost), and .pptx (PowerPoint - each slide's "
                "text, slides separated by a line of just '---')."
            ),
            "parameters": {
                "type": "object",
                "properties": {"path": {"type": "string", "description": "Path of the file to read."}},
                "required": ["path"],
            },
        },
        "confirm_text": lambda a: f"Read the contents of this file?\n\n{a.get('path')}",
        "run": _read_file,
    },
    "write_file": {
        "declaration": {
            "name": "write_file",
            "description": (
                "Create a new file, or overwrite an existing file, with the given text "
                "content. Used to edit documents. For plain text files, content is "
                "written as-is. For .docx, content is written as one paragraph per "
                "line - a brand-new Word document, NOT a preserved edit of the "
                "original's formatting/images. For .xlsx, content must be "
                "comma-separated values, one row per line - a brand-new single-sheet "
                "workbook. For .pdf, content is drawn as plain lines of text on a "
                "fresh document - no original layout, fonts, or images are preserved. "
                "For .pptx, content is one slide per block, separated by a line "
                "containing exactly '---': the block's first line becomes the slide "
                "title, remaining lines become bullet points on that slide - a "
                "brand-new plain deck using PowerPoint's default template, no custom "
                "design/images/animations."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path of the file to create/overwrite."},
                    "content": {"type": "string", "description": "The full new text content of the file."},
                },
                "required": ["path", "content"],
            },
        },
        "confirm_text": lambda a: (
            f"{_write_warning(a.get('path', ''))}"
            f"Overwrite/create this file?\n\n{a.get('path')}\n\n"
            f"--- New full content ---\n{a.get('content', '')[:4000]}"
            + ("\n...(truncated for preview - full content will still be written)" if len(a.get('content', '')) > 4000 else "")
        ),
        "run": _write_file,
    },
    "create_presentation": {
        "declaration": {
            "name": "create_presentation",
            "description": (
                "Create a professionally designed PowerPoint (.pptx) with a dark "
                "teal-on-black theme, multiple slide layouts, native charts, stock "
                "photos from the web, and slide transitions. ALWAYS prefer this over "
                "write_file for any presentation or slide deck request - write_file "
                "only makes plain unstyled title+bullet slides.\n\n"
                "Design rules to follow when writing the spec:\n"
                "- VARY THE LAYOUTS. Never use 'bullets' for every slide. A strong "
                "deck opens with 'title', uses 'section' dividers between themes, and "
                "mixes 'stats', 'compare', 'steps', 'chart', 'image' and 'quote' "
                "throughout.\n"
                "- Every slide should carry a visual. Add an 'image' field (a short "
                "stock-photo search phrase like 'solar panels desert' or 'team "
                "collaboration office') to title/bullets/image slides.\n"
                "- Prefer a 'chart' slide over listing numbers as bullets.\n"
                "- Keep bullets short - 6 words or so each, max 5 per slide. Put the "
                "detail in 'notes' (speaker notes) instead of crowding the slide.\n"
                "- Aim for 8-14 slides for a normal topic unless asked otherwise.\n\n"
                "The 'spec' argument is a JSON string:\n"
                '{"transition": "fade", "slides": [...]}\n'
                "transition (optional, applies to all slides): fade | push | wipe | "
                "morph | split. Any slide may override it with its own 'transition'.\n\n"
                "Every slide takes optional 'notes' (speaker notes). Layouts:\n"
                '- {"layout":"title","title":...,"subtitle":...,"image":...,"footer":...}\n'
                '- {"layout":"section","eyebrow":"PART ONE","title":...}\n'
                '- {"layout":"bullets","title":...,"subtitle":...,"bullets":[...],"image":...}\n'
                '- {"layout":"stats","title":...,"stats":[{"value":"87%","label":"..."}]} (max 4)\n'
                '- {"layout":"compare","title":...,"left":{"heading":...,"points":[...]},'
                '"right":{"heading":...,"points":[...]}}\n'
                '- {"layout":"steps","title":...,"steps":[{"heading":...,"detail":...}]} (max 5)\n'
                '- {"layout":"chart","title":...,"chart":{"type":"bar|hbar|line|pie|area|'
                'doughnut","categories":[...],"series":[{"name":...,"values":[...]}]}}\n'
                '- {"layout":"image","title":...,"image":...,"caption":...}\n'
                '- {"layout":"quote","quote":...,"attribution":...}'
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": "Path of the .pptx file to create.",
                    },
                    "spec": {
                        "type": "string",
                        "description": "The presentation spec as a JSON string (see description).",
                    },
                },
                "required": ["path", "spec"],
            },
        },
        "confirm_text": _presentation_confirm_text,
        "run": _create_presentation,
    },
}

FUNCTION_DECLARATIONS = [a["declaration"] for a in ACTIONS.values()]


def confirm_text_for(name, args):
    if name not in ACTIONS:
        return f"Run unknown action '{name}' with {args}?"
    return ACTIONS[name]["confirm_text"](args)


def run_action(name, args):
    if name not in ACTIONS:
        return {"error": f"Unknown action '{name}'."}
    return ACTIONS[name]["run"](args)
